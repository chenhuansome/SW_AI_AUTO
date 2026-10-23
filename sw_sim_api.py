import os
from docx import Document
from docx.shared import Inches

# ====================== 原有函数保留 ======================
def get_sim_result(sw):
    model = sw.ActiveDoc
    if model is None:
        return {"error": "没有打开零件文档"}
    mass = model.MassProperties(0,0,0,0,0,0,0,0,0,0,0,0,0)[0]
    return {
        "零件质量(kg)": round(mass,3)
    }

def check_result(data, rule_str):
    ok_flag = True
    msg = []
    if ok_flag:
        return "✅全部指标合格"
    else:
        return f"❌不合格，超标项：{msg}"

def export_mode_animation(sw, case_name:str, mode_order:int, save_path:str):
    model = sw.ActiveDoc
    if model is None:
        return "❌ 没有激活零件文档"
    simAddIn = sw.GetAddInObject("SldWorks.SimulationAddIn")
    if simAddIn is None:
        return "❌ 无法获取Simulation插件对象，请确认Simulation插件已经提前勾选加载！"
    simDoc = simAddIn.GetSimulationDoc(model)
    study = simDoc.GetStudyByName(case_name)
    if study is None:
        return f"❌ 找不到模态算例：{case_name}"
    result = study.GetResult()
    result.SelectMode(mode_order)
    anim = result.GetAnimation()
    anim.Play()
    anim.SaveAvi(save_path, 25, True)
    anim.Stop()
    return f"✅振型动画导出成功：{save_path}"


# ====================== 新增1：静力学应力云图动画导出 ======================
def export_static_stress_animation(sw, case_name:str, save_path:str):
    """
    静力学算例：应力云图变形动画导出avi
    case_name: 静力学算例名称，例如 "静应力算例1"
    save_path: avi输出完整路径
    """
    model = sw.ActiveDoc
    if model is None:
        return "❌ 没有激活零件文档"
    simAddIn = sw.GetAddInObject("SldWorks.SimulationAddIn")
    if simAddIn is None:
        return "❌ Simulation插件未加载，获取AddIn对象失败"
    simDoc = simAddIn.GetSimulationDoc(model)
    study = simDoc.GetStudyByName(case_name)
    if study is None:
        return f"❌ 找不到静力学算例：{case_name}"
    result = study.GetResult()
    # 激活应力结果
    stressRes = result.GetStressResult()
    anim = stressRes.GetAnimation()
    anim.Play()
    anim.SaveAvi(save_path, 25, True)
    anim.Stop()
    return f"✅静力学应力变形动画导出成功：{save_path}"


# ====================== 新增2：自动生成Word仿真报告 ======================
def create_simulation_word_report(sw, report_title:str, sim_data:dict, screenshot_path:str, save_docx_path:str):
    """
    report_title: 报告标题
    sim_data: 字典，例如 {"最大应力": "82MPa", "最大位移": "0.35mm", "一阶频率":"215Hz"}
    screenshot_path: 仿真云图截图图片路径
    save_docx_path: 最终docx保存完整路径
    """
    doc = Document()
    doc.add_heading(report_title, level=1)

    doc.add_heading("一、仿真基本参数", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "项目"
    hdr_cells[1].text = "数值"
    for k, v in sim_data.items():
        row_cells = table.add_row().cells
        row_cells[0].text = k
        row_cells[1].text = str(v)

    doc.add_heading("二、仿真云图", level=2)
    if os.path.exists(screenshot_path):
        doc.add_picture(screenshot_path, width=Inches(5.5))
    else:
        doc.add_paragraph("⚠️未找到云图截图文件")

    doc.add_heading("三、结论", level=2)
    doc.add_paragraph("根据仿真结果，对照设计许用值，判定结构是否满足强度、刚度、模态要求。")

    doc.save(save_docx_path)
    return f"✅Word仿真报告生成成功：{save_docx_path}"